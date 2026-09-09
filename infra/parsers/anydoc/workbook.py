"""从原始 Excel 字节补充 AnyDoc 不提供的工作表视图证据。"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from typing import Any


@dataclass(frozen=True)
class WorkbookCellEvidence:
    raw_value: Any
    value_type: str
    formula: str | None = None


@dataclass(frozen=True)
class WorkbookSheetEvidence:
    name: str
    source_row_count: int
    visible_row_count: int
    source_col_count: int
    hidden_row_count: int
    source_has_filter: bool | None
    source_range: str | None
    source_row_numbers: tuple[int, ...]
    visible_source_row_numbers: tuple[int, ...]
    cells: dict[tuple[int, int], WorkbookCellEvidence]
    metadata: dict[str, Any]


def _is_ole2(content: bytes) -> bool:
    return content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _portable_value(value: Any) -> Any:
    """把工作簿值限制为结构 JSON 能稳定表达的原子类型。"""

    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _xlsx_value_type(cell: Any) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.is_date:
        return "date"
    return {
        "s": "string",
        "inlineStr": "string",
        "n": "number",
        "b": "boolean",
        "e": "error",
        "d": "date",
    }.get(cell.data_type, "blank" if cell.value is None else "unknown")


def _xlsx_evidence(content: bytes) -> list[WorkbookSheetEvidence]:
    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=False)
    result: list[WorkbookSheetEvidence] = []
    for sheet in workbook.worksheets:
        source_rows: list[int] = []
        visible_source_rows: list[int] = []
        max_col = 0
        cells: dict[tuple[int, int], WorkbookCellEvidence] = {}
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            occupied = [
                index for index, cell in enumerate(row, start=1)
                if cell.value not in (None, "")
            ]
            if not occupied:
                continue
            source_rows.append(row_number)
            max_col = max(max_col, occupied[-1])
            if not bool(sheet.row_dimensions[row_number].hidden):
                visible_source_rows.append(row_number)
            for col_number, cell in enumerate(row, start=1):
                if col_number > max_col:
                    break
                value = _portable_value(cell.value)
                cells[(row_number, col_number)] = WorkbookCellEvidence(
                    raw_value=value,
                    value_type=_xlsx_value_type(cell),
                    formula=str(cell.value) if cell.data_type == "f" else None,
                )
        if not source_rows:
            continue
        filter_ref = getattr(sheet.auto_filter, "ref", None)
        result.append(
            WorkbookSheetEvidence(
                name=sheet.title,
                source_row_count=len(source_rows),
                visible_row_count=len(visible_source_rows),
                source_col_count=max_col,
                hidden_row_count=len(source_rows) - len(visible_source_rows),
                source_has_filter=bool(filter_ref),
                source_range=sheet.calculate_dimension() if max_col else None,
                source_row_numbers=tuple(source_rows),
                visible_source_row_numbers=tuple(visible_source_rows),
                cells=cells,
                metadata={
                    "filter_ref": filter_ref,
                    "has_hidden_rows": len(visible_source_rows) != len(source_rows),
                    "merged_range_count": len(sheet.merged_cells.ranges),
                    "workbook_reader": "openpyxl",
                },
            )
        )
    workbook.close()
    return result


def _xls_evidence(content: bytes) -> list[WorkbookSheetEvidence]:
    import xlrd
    from openpyxl.utils import get_column_letter

    workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    result: list[WorkbookSheetEvidence] = []
    for sheet in workbook.sheets():
        source_rows: list[int] = []
        visible_source_rows: list[int] = []
        cells: dict[tuple[int, int], WorkbookCellEvidence] = {}
        max_col = 0
        min_row: int | None = None
        min_col: int | None = None
        for row_index in range(sheet.nrows):
            occupied = [
                col_index + 1
                for col_index in range(sheet.ncols)
                if sheet.cell_value(row_index, col_index) not in (None, "")
            ]
            if not occupied:
                continue
            source_row = row_index + 1
            source_rows.append(source_row)
            max_col = max(max_col, occupied[-1])
            min_row = source_row if min_row is None else min(min_row, source_row)
            min_col = min(occupied) if min_col is None else min(min_col, min(occupied))
            row_info = sheet.rowinfo_map.get(row_index)
            if row_info is None or not row_info.hidden:
                visible_source_rows.append(source_row)
            for col_index in range(max_col):
                cell = sheet.cell(row_index, col_index)
                type_name = {
                    xlrd.XL_CELL_EMPTY: "blank",
                    xlrd.XL_CELL_TEXT: "string",
                    xlrd.XL_CELL_NUMBER: "number",
                    xlrd.XL_CELL_DATE: "date_serial",
                    xlrd.XL_CELL_BOOLEAN: "boolean",
                    xlrd.XL_CELL_ERROR: "error",
                    xlrd.XL_CELL_BLANK: "blank",
                }.get(cell.ctype, "unknown")
                cells[(source_row, col_index + 1)] = WorkbookCellEvidence(
                    raw_value=_portable_value(cell.value),
                    value_type=type_name,
                )
        if not source_rows:
            continue
        source_range = (
            f"{get_column_letter(min_col or 1)}{min_row or 1}:"
            f"{get_column_letter(max_col)}{max(source_rows)}"
            if max_col
            else None
        )
        result.append(
            WorkbookSheetEvidence(
                name=sheet.name,
                source_row_count=len(source_rows),
                visible_row_count=len(visible_source_rows),
                source_col_count=max_col,
                hidden_row_count=len(source_rows) - len(visible_source_rows),
                # xlrd 不提供足够证据区分 AutoFilter 和手工隐藏。
                source_has_filter=None,
                source_range=source_range,
                source_row_numbers=tuple(source_rows),
                visible_source_row_numbers=tuple(visible_source_rows),
                cells=cells,
                metadata={
                    "has_hidden_rows": len(visible_source_rows) != len(source_rows),
                    "merged_range_count": len(sheet.merged_cells),
                    "workbook_reader": "xlrd",
                },
            )
        )
    return result


def inspect_workbook(content: bytes, extension: str) -> list[WorkbookSheetEvidence]:
    """读取工作表计数、隐藏行和合并数量；失败时由调用方安全降级。"""

    if _is_ole2(content):
        return _xls_evidence(content)
    if extension.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _xlsx_evidence(content)
    return []


def enrich_tables_with_workbook(
    payload: dict[str, Any],
    *,
    content: bytes,
    extension: str,
) -> tuple[dict[str, Any], list[str]]:
    """将工作簿级证据合入通用表格 payload，不伪造源单元格地址。"""

    warnings: list[str] = []
    from openpyxl.utils import get_column_letter
    try:
        sheets = inspect_workbook(content, extension)
    except Exception as error:
        return payload, [f"工作簿结构预处理失败：{type(error).__name__}: {error}"]
    tables = payload.get("tables")
    if not sheets or not isinstance(tables, list):
        return payload, warnings

    unused = list(sheets)
    enriched: list[Any] = []
    for table in tables:
        if not isinstance(table, dict):
            enriched.append(table)
            continue
        name = table.get("source_container_name")
        evidence = next((sheet for sheet in unused if sheet.name == name), None)
        if evidence is None:
            rows = table.get("num_rows")
            cols = table.get("num_cols")
            candidates = [
                sheet
                for sheet in unused
                if sheet.source_col_count == cols
                and rows in {sheet.source_row_count, sheet.visible_row_count}
            ]
            if len(candidates) == 1:
                evidence = candidates[0]
        if evidence is None and len(unused) == 1:
            evidence = unused[0]
        if evidence is None:
            enriched.append(table)
            warnings.append(f"table {table.get('table_id')} 无法唯一映射到源工作表。")
            continue
        unused.remove(evidence)
        emitted = table.get("num_rows") if isinstance(table.get("num_rows"), int) else None
        if emitted == evidence.visible_row_count < evidence.source_row_count:
            view_scope = "visible_rows"
            source_rows = evidence.visible_source_row_numbers
        elif emitted == evidence.source_row_count:
            view_scope = "all_rows"
            source_rows = evidence.source_row_numbers
        else:
            view_scope = "unknown"
            source_rows = ()
        raw_cells = table.get("cells")
        enriched_cells: list[Any] = []
        can_map_columns = table.get("num_cols") == evidence.source_col_count
        if isinstance(raw_cells, list):
            for raw_cell in raw_cells:
                if not isinstance(raw_cell, dict):
                    enriched_cells.append(raw_cell)
                    continue
                logical_row = raw_cell.get("start_row")
                logical_col = raw_cell.get("start_col")
                if (
                    not source_rows
                    or not can_map_columns
                    or not isinstance(logical_row, int)
                    or not isinstance(logical_col, int)
                    or logical_row < 0
                    or logical_row >= len(source_rows)
                    or logical_col < 0
                ):
                    enriched_cells.append(raw_cell)
                    continue
                source_row = source_rows[logical_row]
                source_col = logical_col + 1
                source_cell = evidence.cells.get((source_row, source_col))
                cell_ref = f"{get_column_letter(source_col)}{source_row}"
                anchor = (
                    raw_cell.get("source_anchor")
                    if isinstance(raw_cell.get("source_anchor"), dict)
                    else {}
                )
                additions: dict[str, Any] = {
                    "source_anchor": {
                        **anchor,
                        "container": "sheet",
                        "container_name": evidence.name,
                        "cell_ref": cell_ref,
                        "range_ref": evidence.source_range,
                        "table_cell": f"r{logical_row}c{logical_col}",
                        "provenance_status": "available",
                    },
                    "visible": source_row in evidence.visible_source_row_numbers,
                }
                if source_cell is not None:
                    additions.update(
                        {
                            "raw_value": source_cell.raw_value,
                            "value_type": source_cell.value_type,
                            "formula": source_cell.formula,
                        }
                    )
                enriched_cells.append({**raw_cell, **additions})
        if source_rows and not can_map_columns:
            warnings.append(
                f"table {table.get('table_id')} 列数与源工作表不一致，未生成单元格 A1 定位。"
            )
        enriched.append(
            {
                **table,
                "source_container": "sheet",
                "source_container_name": evidence.name,
                "source_range": evidence.source_range,
                "view_scope": view_scope,
                "source_has_filter": evidence.source_has_filter,
                "source_row_count": evidence.source_row_count,
                "emitted_row_count": emitted,
                "hidden_row_count": evidence.hidden_row_count,
                "cells": enriched_cells if isinstance(raw_cells, list) else raw_cells,
                "metadata": {
                    **(table.get("metadata") if isinstance(table.get("metadata"), dict) else {}),
                    **evidence.metadata,
                    "source_col_count": evidence.source_col_count,
                },
            }
        )
    return {**payload, "tables": enriched}, warnings
