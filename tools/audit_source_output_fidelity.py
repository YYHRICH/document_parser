"""Build a reproducible source-to-parser-output fidelity baseline.

The source corpus and saved parser results are strictly read-only. This tool writes
compact JSONL/JSON artifacts only under the project workspace.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import statistics
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path, PureWindowsPath
from typing import Any, Iterable
from xml.etree import ElementTree

import olefile
from openpyxl import load_workbook
from pypdf import PdfReader


CATEGORIES = (
    "图文回答数据集源文件",
    "多文档多段知识数据集源文件",
    "文档单点知识数据集源文件",
    "模糊回答数据集源文件",
)
MODEL_DIRS = {
    "anydoc": "res_anydoc",
    "docling": "res_docling",
    "markitdown": "res_markitdown",
    "mineru": "res_minerU",
}
SUCCESS_STATES = {"success", "reused", "done", "completed"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".svg"}
XML_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--results-root", type=Path, required=True)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("artifacts") / "source-output-fidelity",
    )
    parser.add_argument("--sample-count", type=int, default=28)
    parser.add_argument(
        "--reuse-document-metrics",
        action="store_true",
        help="Regenerate sample/summary artifacts from existing document_metrics.jsonl.",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_key(*parts: str) -> str:
    return hashlib.sha256("/".join(parts).encode("utf-8")).hexdigest()[:24]


def read_text(path: Path | None) -> str:
    if path is None or not path.is_file():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def read_json(path: Path | None, default: Any) -> Any:
    if path is None or not path.is_file():
        return default
    try:
        return json.loads(read_text(path))
    except (json.JSONDecodeError, OSError):
        return default


def write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True))
            handle.write("\n")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [
        json.loads(line)
        for line in read_text(path).splitlines()
        if line.strip()
    ]


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).lower()
    return "".join(char for char in value if char.isalnum() or "\u4e00" <= char <= "\u9fff")


def relative_display(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return str(path)


def first_file(directory: Path, pattern: str) -> Path | None:
    if not directory.is_dir():
        return None
    try:
        return next(iter(sorted(directory.glob(pattern), key=lambda item: item.name)), None)
    except OSError:
        return None


def windows_name(value: str) -> str:
    return PureWindowsPath(value).name


def canonical_windows_component(value: str) -> str:
    return unicodedata.normalize("NFKC", value).rstrip(" .").casefold()


def resolve_result_directory(
    *,
    results_root: Path,
    model_dir: str,
    category: str,
    filename: str,
    manifest_item: dict[str, Any] | None = None,
) -> Path:
    category_dir = results_root / model_dir / category
    requested_names = [Path(filename).stem]
    if manifest_item:
        output_dir = manifest_item.get("output_dir")
        if output_dir:
            requested_names.append(windows_name(str(output_dir)))
        markdown = manifest_item.get("markdown")
        if markdown:
            markdown_path = PureWindowsPath(str(markdown))
            if len(markdown_path.parts) > 1 and not markdown_path.is_absolute():
                requested_names.append(markdown_path.parts[0])
    targets = {canonical_windows_component(name) for name in requested_names}
    matches = [
        path
        for path in category_dir.iterdir()
        if path.is_dir() and canonical_windows_component(path.name) in targets
    ] if category_dir.is_dir() else []
    if len(matches) == 1:
        return matches[0]
    for name in requested_names:
        candidate = category_dir / name
        if candidate.is_dir():
            return candidate
    return category_dir / Path(filename).stem


def page_image_count(page: Any) -> int:
    try:
        resources = page.get("/Resources") or {}
        xobjects = resources.get("/XObject") or {}
        xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
        count = 0
        for item in xobjects.values():
            obj = item.get_object() if hasattr(item, "get_object") else item
            if obj.get("/Subtype") == "/Image":
                count += 1
        return count
    except Exception:
        return 0


def preflight_pdf(path: Path) -> tuple[dict[str, Any], str]:
    result: dict[str, Any] = {
        "source_kind": "pdf",
        "page_count": None,
        "sheet_count": None,
        "has_text_layer": None,
        "scanned_page_ratio": None,
        "text_page_count": 0,
        "low_text_page_count": 0,
        "source_text_chars": 0,
        "source_image_object_count": 0,
        "source_table_count": None,
        "source_media_count": None,
        "is_corrupt": False,
        "is_encrypted": False,
        "preflight_warnings": [],
        "text_layer_estimate": "pypdf extract_text; low-text means below 20 normalized characters",
    }
    text_parts: list[str] = []
    try:
        reader = PdfReader(str(path), strict=False)
        result["is_encrypted"] = bool(reader.is_encrypted)
        result["page_count"] = len(reader.pages)
        for page_index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as error:
                text = ""
                result["preflight_warnings"].append(
                    f"page {page_index} text extraction failed: {type(error).__name__}"
                )
            normalized_length = len(normalize_text(text))
            if normalized_length >= 20:
                result["text_page_count"] += 1
            else:
                result["low_text_page_count"] += 1
            text_parts.append(text)
            result["source_image_object_count"] += page_image_count(page)
        page_count = result["page_count"] or 0
        result["has_text_layer"] = result["text_page_count"] > 0
        result["scanned_page_ratio"] = (
            round(result["low_text_page_count"] / page_count, 6) if page_count else None
        )
    except Exception as error:
        result["is_corrupt"] = True
        result["preflight_warnings"].append(
            f"PDF open failed: {type(error).__name__}: {str(error)[:200]}"
        )
    source_text = "\n".join(text_parts)
    result["source_text_chars"] = len(source_text)
    return result, source_text


def preflight_docx(path: Path) -> tuple[dict[str, Any], str]:
    result: dict[str, Any] = {
        "source_kind": "docx",
        "page_count": None,
        "sheet_count": None,
        "has_text_layer": True,
        "scanned_page_ratio": None,
        "source_text_chars": 0,
        "source_table_count": 0,
        "source_media_count": 0,
        "source_image_object_count": None,
        "is_corrupt": False,
        "preflight_warnings": [],
    }
    text = ""
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            document = ElementTree.fromstring(archive.read("word/document.xml"))
            text = "\n".join(
                node.text or "" for node in document.findall(".//w:t", XML_NS)
            )
            result["source_table_count"] = len(document.findall(".//w:tbl", XML_NS))
            result["source_media_count"] = sum(
                name.startswith("word/media/") and not name.endswith("/") for name in names
            )
            if "docProps/app.xml" in names:
                app = ElementTree.fromstring(archive.read("docProps/app.xml"))
                pages = app.find("ep:Pages", XML_NS)
                if pages is not None and (pages.text or "").isdigit():
                    result["page_count"] = int(pages.text or "0")
    except Exception as error:
        result["is_corrupt"] = True
        result["preflight_warnings"].append(
            f"DOCX open failed: {type(error).__name__}: {str(error)[:200]}"
        )
    result["source_text_chars"] = len(text)
    return result, text


def preflight_xlsx(path: Path) -> tuple[dict[str, Any], str]:
    result: dict[str, Any] = {
        "source_kind": "xlsx",
        "page_count": None,
        "sheet_count": None,
        "has_text_layer": True,
        "scanned_page_ratio": None,
        "source_text_chars": 0,
        "source_table_count": 0,
        "source_cell_count": 0,
        "source_media_count": 0,
        "source_image_object_count": None,
        "is_corrupt": False,
        "preflight_warnings": [],
    }
    parts: list[str] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        result["sheet_count"] = len(workbook.worksheets)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        result["source_cell_count"] += 1
                        parts.append(str(cell.value))
        workbook.close()
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            result["source_table_count"] = sum(
                name.startswith("xl/tables/") and name.endswith(".xml") for name in names
            )
            result["source_media_count"] = sum(
                name.startswith("xl/media/") and not name.endswith("/") for name in names
            )
    except Exception as error:
        result["is_corrupt"] = True
        result["preflight_warnings"].append(
            f"XLSX open failed: {type(error).__name__}: {str(error)[:200]}"
        )
    text = "\n".join(parts)
    result["source_text_chars"] = len(text)
    return result, text


def preflight_legacy_office(path: Path) -> tuple[dict[str, Any], str]:
    result: dict[str, Any] = {
        "source_kind": path.suffix.lower().lstrip("."),
        "page_count": None,
        "sheet_count": None,
        "has_text_layer": None,
        "scanned_page_ratio": None,
        "source_text_chars": None,
        "source_table_count": None,
        "source_media_count": None,
        "source_image_object_count": None,
        "is_corrupt": False,
        "preflight_warnings": [
            "Legacy OLE preflight exposes container metadata only; sampled inspection is required."
        ],
    }
    try:
        ole = olefile.OleFileIO(str(path))
        metadata = ole.get_metadata()
        pages = getattr(metadata, "num_pages", None)
        if isinstance(pages, int) and pages > 0:
            result["page_count"] = pages
        stream_names = ["/".join(parts) for parts in ole.listdir()]
        result["ole_stream_count"] = len(stream_names)
        result["ole_stream_names"] = stream_names[:30]
        ole.close()
    except Exception as error:
        result["is_corrupt"] = True
        result["preflight_warnings"].append(
            f"OLE open failed: {type(error).__name__}: {str(error)[:200]}"
        )
    return result, ""


def source_preflight(path: Path) -> tuple[dict[str, Any], str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return preflight_pdf(path)
    if suffix == ".docx":
        return preflight_docx(path)
    if suffix == ".xlsx":
        return preflight_xlsx(path)
    if suffix in {".doc", ".ppt", ".xls"}:
        return preflight_legacy_office(path)
    return (
        {
            "source_kind": suffix.lstrip("."),
            "page_count": None,
            "sheet_count": None,
            "has_text_layer": None,
            "scanned_page_ratio": None,
            "source_text_chars": None,
            "source_table_count": None,
            "source_media_count": None,
            "source_image_object_count": None,
            "is_corrupt": False,
            "preflight_warnings": ["No format-specific preflight implemented."],
        },
        "",
    )


def markdown_blocks(markdown: str) -> list[str]:
    blocks: list[str] = []
    pending: list[str] = []
    in_fence = False
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if re.match(r"^(```+|~~~+)", line):
            in_fence = not in_fence
        if not in_fence and line.startswith("#"):
            if pending:
                blocks.append("\n".join(pending).strip())
                pending.clear()
            blocks.append(line)
        elif not line and not in_fence:
            if pending:
                blocks.append("\n".join(pending).strip())
                pending.clear()
        else:
            pending.append(raw_line)
    if pending:
        blocks.append("\n".join(pending).strip())
    return [block for block in blocks if block]


def markdown_metrics(markdown_paths: list[Path], asset_root: Path) -> tuple[dict[str, Any], str]:
    combined_parts: list[str] = []
    image_refs: list[tuple[Path, str]] = []
    for path in markdown_paths:
        text = read_text(path)
        combined_parts.append(text)
        for match in re.finditer(r"!\[[^\]]*\]\(([^\)]+)\)", text):
            image_refs.append((path, match.group(1).strip().strip("<>")))
    markdown = "\n\n".join(combined_parts)
    blocks = markdown_blocks(markdown)
    normalized_blocks = [normalize_text(block) for block in blocks]
    counts = Counter(block for block in normalized_blocks if block)
    duplicate_block_count = sum(count - 1 for count in counts.values() if count > 1)
    max_block_chars = max((len(block) for block in blocks), default=0)
    missing_refs = 0
    data_uri_count = 0
    external_refs = 0
    for markdown_path, reference in image_refs:
        lowered = reference.lower()
        if lowered.startswith("data:image/"):
            data_uri_count += 1
            continue
        if lowered.startswith(("http://", "https://")):
            external_refs += 1
            continue
        clean_reference = reference.split("#", 1)[0].split("?", 1)[0]
        if not (markdown_path.parent / clean_reference).is_file():
            missing_refs += 1
    asset_files = [
        path
        for path in asset_root.rglob("*")
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    ]
    heading_count = len(re.findall(r"(?m)^#{1,6}\s+\S", markdown))
    html_table_count = len(re.findall(r"(?is)<table\b", markdown))
    pipe_table_row_count = len(re.findall(r"(?m)^\s*\|.*\|\s*$", markdown))
    total_chars = len(markdown)
    return (
        {
            "markdown_file_count": len(markdown_paths),
            "markdown_chars": total_chars,
            "markdown_lines": len(markdown.splitlines()),
            "markdown_sha256": hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
            "markdown_block_count": len(blocks),
            "heading_count": heading_count,
            "html_table_count": html_table_count,
            "pipe_table_row_count": pipe_table_row_count,
            "image_reference_count": len(image_refs),
            "data_uri_count": data_uri_count,
            "external_image_reference_count": external_refs,
            "missing_image_reference_count": missing_refs,
            "actual_asset_count": len(asset_files),
            "duplicate_block_count": duplicate_block_count,
            "duplicate_block_ratio": (
                round(duplicate_block_count / len(blocks), 6) if blocks else 0.0
            ),
            "max_block_chars": max_block_chars,
            "largest_block_share": (
                round(max_block_chars / total_chars, 6) if total_chars else 0.0
            ),
            "has_giant_block": bool(
                max_block_chars >= 10000
                or (total_chars >= 1000 and max_block_chars / total_chars >= 0.8)
            ),
            "is_empty": not bool(markdown.strip()),
        },
        markdown,
    )


def text_fidelity_proxy(source_text: str, output_markdown: str) -> dict[str, Any]:
    source = normalize_text(source_text)
    output = normalize_text(output_markdown)
    if not source:
        return {
            "available": False,
            "reason": "source preflight did not produce reference text",
            "source_normalized_chars": 0,
            "output_normalized_chars": len(output),
        }
    if not output:
        return {
            "available": True,
            "source_normalized_chars": len(source),
            "output_normalized_chars": 0,
            "anchor_count": 0,
            "matched_anchor_count": 0,
            "anchor_coverage": 0.0,
            "anchor_order_score": 0.0,
            "length_ratio": 0.0,
        }
    anchor_width = 24
    desired_anchors = min(200, max(20, len(source) // 500))
    step = max(anchor_width, len(source) // desired_anchors)
    anchors = [
        (position, source[position : position + anchor_width])
        for position in range(0, max(1, len(source) - anchor_width + 1), step)
        if len(source[position : position + anchor_width]) >= 12
    ]
    matches: list[tuple[int, int]] = []
    for source_position, anchor in anchors:
        output_position = output.find(anchor)
        if output_position >= 0:
            matches.append((source_position, output_position))
    monotonic = sum(
        current[1] >= previous[1]
        for previous, current in zip(matches, matches[1:])
    )
    return {
        "available": True,
        "method": "exact normalized 24-character anchors sampled across source text",
        "source_normalized_chars": len(source),
        "output_normalized_chars": len(output),
        "anchor_count": len(anchors),
        "matched_anchor_count": len(matches),
        "anchor_coverage": round(len(matches) / len(anchors), 6) if anchors else None,
        "anchor_order_score": (
            round(monotonic / (len(matches) - 1), 6) if len(matches) > 1 else None
        ),
        "length_ratio": round(len(output) / len(source), 6),
    }


def manifest_index(results_root: Path, model_dir: str) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for category in CATEGORIES:
        manifest = read_json(results_root / model_dir / category / "manifest.json", {})
        for item in manifest.get("files", []):
            source_name = windows_name(str(item.get("source") or ""))
            index[f"{category}/{source_name}".lower()] = item
    return index


def base_output_record(parser_id: str, status: str) -> dict[str, Any]:
    return {
        "parser": parser_id,
        "parser_status": status,
        "parser_success": status in SUCCESS_STATES,
        "result_exists": False,
        "error_code": None,
        "error": None,
        "markdown_file_count": 0,
        "markdown_chars": 0,
        "markdown_block_count": 0,
        "heading_count": 0,
        "native_block_count": 0,
        "native_table_count": 0,
        "native_table_cell_count": 0,
        "native_picture_count": 0,
        "native_page_coverage_count": 0,
        "page_coverage_ratio": None,
        "native_bbox_coverage": None,
        "native_source_id_coverage": None,
        "normalized_source_id_coverage": None,
        "source_id_origin": None,
        "root_structure_merged": None,
        "split_document": False,
        "parts_expected": None,
        "parts_actual": None,
        "parts_complete": None,
        "page_ranges_contiguous": None,
        "text_fidelity_proxy": {"available": False, "reason": "no output"},
    }


def manifest_output_metrics(
    *,
    parser_id: str,
    model_dir: str,
    source_record: dict[str, Any],
    source_text: str,
    results_root: Path,
    index: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    key = f"{source_record['category']}/{source_record['filename']}".lower()
    item = index.get(key, {})
    status = str(item.get("status") or "missing").lower()
    result = base_output_record(parser_id, status)
    result["error_code"] = item.get("error_code")
    error = item.get("error")
    result["error"] = re.sub(r"\s+", " ", str(error)).strip()[:500] if error else None
    result["manifest_metrics"] = {
        field: item[field]
        for field in (
            "characters",
            "bytes",
            "text_count",
            "table_count",
            "picture_count",
            "image_file_count",
            "elapsed_ms",
        )
        if field in item
    }
    document_dir = resolve_result_directory(
        results_root=results_root,
        model_dir=model_dir,
        category=source_record["category"],
        filename=source_record["filename"],
        manifest_item=item,
    )
    result["result_directory"] = relative_display(document_dir, results_root)
    if status not in SUCCESS_STATES:
        return result
    markdown_path = first_file(document_dir, "*.md")
    json_path = first_file(document_dir, "*.json")
    markdown_paths = [markdown_path] if markdown_path else []
    metrics, markdown = markdown_metrics(markdown_paths, document_dir)
    result.update(metrics)
    result["result_exists"] = bool(markdown_path or json_path)
    result["markdown_paths"] = [
        relative_display(path, results_root) for path in markdown_paths
    ]
    result["text_fidelity_proxy"] = text_fidelity_proxy(source_text, markdown)
    if parser_id in {"anydoc", "markitdown"}:
        result["native_block_count"] = 0
        result["native_source_id_coverage"] = 0.0
        result["normalized_source_id_coverage"] = 0.0 if metrics["markdown_block_count"] else None
        result["source_id_origin"] = "unavailable_in_current_markdown_builder"
    return result


def docling_provenance(item: dict[str, Any]) -> list[dict[str, Any]]:
    prov = item.get("prov")
    return [entry for entry in prov if isinstance(entry, dict)] if isinstance(prov, list) else []


def docling_page_number(entry: dict[str, Any]) -> int | None:
    value = entry.get("page_no")
    if not isinstance(value, int):
        return None
    return value + 1 if value == 0 else value


def docling_output_metrics(
    *,
    source_record: dict[str, Any],
    source_text: str,
    results_root: Path,
    index: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    result = manifest_output_metrics(
        parser_id="docling",
        model_dir=MODEL_DIRS["docling"],
        source_record=source_record,
        source_text=source_text,
        results_root=results_root,
        index=index,
    )
    if not result["parser_success"]:
        return result
    document_dir = results_root / result["result_directory"]
    json_path = first_file(document_dir, "*.json")
    payload = read_json(json_path, {})
    texts = [item for item in payload.get("texts", []) if isinstance(item, dict)]
    tables = [item for item in payload.get("tables", []) if isinstance(item, dict)]
    pictures = [item for item in payload.get("pictures", []) if isinstance(item, dict)]
    elements = [*texts, *tables, *pictures]
    nonempty_texts = [
        item for item in texts if str(item.get("text") or item.get("orig") or "").strip()
    ]
    pages: set[int] = set()
    elements_with_bbox = 0
    elements_with_source_id = 0
    for item in elements:
        if item.get("self_ref"):
            elements_with_source_id += 1
        item_has_bbox = False
        for entry in docling_provenance(item):
            page_number = docling_page_number(entry)
            if page_number is not None:
                pages.add(page_number)
            if isinstance(entry.get("bbox"), dict):
                item_has_bbox = True
        if item_has_bbox:
            elements_with_bbox += 1
    table_cells = 0
    for table in tables:
        data = table.get("data") if isinstance(table.get("data"), dict) else {}
        cells = data.get("table_cells")
        if isinstance(cells, list):
            table_cells += len(cells)
    source_pages = (
        source_record["preflight"].get("page_count")
        if source_record["extension"] == ".pdf"
        else None
    )
    result.update(
        {
            "native_block_count": len(nonempty_texts) + len(tables),
            "native_text_count": len(texts),
            "native_table_count": len(tables),
            "native_table_cell_count": table_cells,
            "native_picture_count": len(pictures),
            "native_page_coverage_count": len(pages),
            "native_page_numbers": sorted(pages),
            "page_coverage_ratio": (
                round(len(pages) / source_pages, 6)
                if isinstance(source_pages, int) and source_pages > 0
                else None
            ),
            "native_bbox_coverage": (
                round(elements_with_bbox / len(elements), 6) if elements else None
            ),
            "native_source_id_coverage": (
                round(elements_with_source_id / len(elements), 6) if elements else None
            ),
            "normalized_source_id_coverage": None,
            "source_id_origin": (
                "raw Docling self_ref measured; normalized adapter output was not persisted"
            ),
            "native_json_path": (
                relative_display(json_path, results_root) if json_path else None
            ),
        }
    )
    return result


def mineru_content_lists(directory: Path, split_document: bool) -> list[Path]:
    candidates = (
        directory.glob("parts/*/*_content_list.json")
        if split_document
        else directory.glob("*_content_list.json")
    )
    return sorted(
        (
            path
            for path in candidates
            if not path.name.endswith("_content_list_v2.json")
        ),
        key=lambda path: path.as_posix(),
    )


def mineru_part_range(path: Path) -> tuple[int, int] | None:
    for parent in (path.parent, *path.parents):
        match = re.fullmatch(r"part_(\d{4})-(\d{4})", parent.name)
        if match:
            return int(match.group(1)), int(match.group(2))
    return None


def mineru_page_number(item: dict[str, Any], content_path: Path) -> int | None:
    page_index = item.get("page_idx")
    if not isinstance(page_index, int):
        return None
    page_range = mineru_part_range(content_path)
    if page_range:
        return page_range[0] + page_index
    return page_index + 1


def mineru_cell_count(html: str) -> int:
    return len(re.findall(r"(?is)<t[dh]\b", html))


def mineru_output_metrics(
    *,
    source_record: dict[str, Any],
    source_text: str,
    results_root: Path,
) -> dict[str, Any]:
    document_dir = resolve_result_directory(
        results_root=results_root,
        model_dir=MODEL_DIRS["mineru"],
        category=source_record["category"],
        filename=source_record["filename"],
    )
    wrapper_path = document_dir / "_mineru_result.json"
    wrapper = read_json(wrapper_path, {})
    status = str(wrapper.get("state") or ("done" if document_dir.is_dir() else "missing")).lower()
    result = base_output_record("mineru", status)
    result["result_directory"] = relative_display(document_dir, results_root)
    result["result_exists"] = document_dir.is_dir()
    result["error"] = str(wrapper.get("err_msg") or "").strip() or None
    if status not in SUCCESS_STATES or not document_dir.is_dir():
        return result
    split_document = bool(
        wrapper.get("oversized_pdf_split") or (document_dir / "parts").is_dir()
    )
    if split_document:
        markdown_paths = sorted(
            (document_dir / "parts").glob("*/full.md"),
            key=lambda path: path.as_posix(),
        )
    else:
        root_markdown = document_dir / "full.md"
        markdown_paths = [root_markdown] if root_markdown.is_file() else []
    content_paths = mineru_content_lists(document_dir, split_document)
    metrics, markdown = markdown_metrics(markdown_paths, document_dir)
    result.update(metrics)
    result["markdown_paths"] = [
        relative_display(path, results_root) for path in markdown_paths
    ]
    result["text_fidelity_proxy"] = text_fidelity_proxy(source_text, markdown)
    result["split_document"] = split_document
    expected_parts = wrapper.get("parts") if isinstance(wrapper.get("parts"), list) else []
    actual_part_dirs = (
        sorted((document_dir / "parts").glob("part_*")) if split_document else []
    )
    result["parts_expected"] = len(expected_parts) if split_document else 1
    result["parts_actual"] = len(actual_part_dirs) if split_document else 1
    result["parts_complete"] = (
        len(expected_parts) == len(actual_part_dirs) == len(content_paths)
        if split_document
        else bool(content_paths)
    )
    ranges = [
        page_range
        for path in actual_part_dirs
        if (page_range := mineru_part_range(path)) is not None
    ]
    ranges.sort()
    result["page_ranges"] = [
        {"start_page": start, "end_page": end} for start, end in ranges
    ]
    result["page_ranges_contiguous"] = (
        all(current[0] == previous[1] + 1 for previous, current in zip(ranges, ranges[1:]))
        if split_document and ranges
        else (True if not split_document else False)
    )
    result["root_structure_merged"] = not split_document and bool(content_paths)
    result["root_markdown_is_index"] = split_document and (document_dir / "full.md").is_file()
    items: list[tuple[dict[str, Any], Path]] = []
    for content_path in content_paths:
        payload = read_json(content_path, [])
        if isinstance(payload, dict):
            payload = payload.get("content_list") or payload.get("items") or []
        if isinstance(payload, list):
            items.extend(
                (item, content_path) for item in payload if isinstance(item, dict)
            )
    pages: set[int] = set()
    bbox_count = 0
    raw_source_id_count = 0
    block_count = 0
    table_count = 0
    table_cells = 0
    picture_count = 0
    type_counts: Counter[str] = Counter()
    for item, content_path in items:
        item_type = str(item.get("type") or "unknown")
        type_counts[item_type] += 1
        page_number = mineru_page_number(item, content_path)
        if page_number is not None:
            pages.add(page_number)
        if isinstance(item.get("bbox"), list) and len(item["bbox"]) == 4:
            bbox_count += 1
        if item.get("id") or item.get("source_block_id"):
            raw_source_id_count += 1
        has_content = bool(
            str(item.get("text") or "").strip()
            or item_type in {"table", "image", "chart", "equation", "formula"}
        )
        if has_content:
            block_count += 1
        if item_type == "table":
            table_count += 1
            html = item.get("table_body") or item.get("html") or ""
            if isinstance(html, str):
                table_cells += mineru_cell_count(html)
        if item_type in {"image", "chart"}:
            picture_count += 1
    source_pages = (
        source_record["preflight"].get("page_count")
        if source_record["extension"] == ".pdf"
        else None
    )
    result.update(
        {
            "native_block_count": block_count,
            "native_item_count": len(items),
            "native_type_counts": dict(type_counts),
            "native_table_count": table_count,
            "native_table_cell_count": table_cells,
            "native_picture_count": picture_count,
            "native_page_coverage_count": len(pages),
            "native_page_numbers": sorted(pages),
            "page_coverage_ratio": (
                round(len(pages) / source_pages, 6)
                if isinstance(source_pages, int) and source_pages > 0
                else None
            ),
            "native_bbox_coverage": (
                round(bbox_count / len(items), 6) if items else None
            ),
            "native_source_id_coverage": (
                round(raw_source_id_count / len(items), 6) if items else None
            ),
            "normalized_source_id_coverage": None,
            "source_id_origin": (
                "raw MinerU id fields measured; normalized adapter output was not persisted"
            ),
            "native_content_paths": [
                relative_display(path, results_root) for path in content_paths
            ],
        }
    )
    return result


def build_source_records(
    source_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    records: list[dict[str, Any]] = []
    source_text_by_key: dict[str, str] = {}
    for category in CATEGORIES:
        category_dir = source_root / category
        if not category_dir.is_dir():
            raise FileNotFoundError(f"Missing source category: {category_dir}")
        for path in sorted(
            (item for item in category_dir.iterdir() if item.is_file()),
            key=lambda item: item.name.casefold(),
        ):
            document_key = f"{category}/{path.name}"
            preflight, source_text = source_preflight(path)
            record = {
                "document_key": document_key,
                "document_id": stable_key(category, path.name),
                "category": category,
                "filename": path.name,
                "extension": path.suffix.lower(),
                "source_path": relative_display(path, source_root),
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "preflight": preflight,
            }
            records.append(record)
            source_text_by_key[document_key] = source_text

    by_hash: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_hash[record["sha256"]].append(record)
    for digest, occurrences in by_hash.items():
        occurrence_keys = sorted(item["document_key"] for item in occurrences)
        categories = sorted({item["category"] for item in occurrences})
        for record in occurrences:
            record["unique_source_id"] = f"sha256:{digest}"
            record["duplicate_occurrence_count"] = len(occurrences)
            record["duplicate_categories"] = categories
            record["duplicate_document_keys"] = occurrence_keys
    return records, source_text_by_key


def attach_outputs(
    records: list[dict[str, Any]],
    source_text_by_key: dict[str, str],
    results_root: Path,
) -> list[dict[str, Any]]:
    indexes = {
        model: manifest_index(results_root, MODEL_DIRS[model])
        for model in ("anydoc", "docling", "markitdown")
    }
    for record in records:
        source_text = source_text_by_key[record["document_key"]]
        record["outputs"] = {
            "anydoc": manifest_output_metrics(
                parser_id="anydoc",
                model_dir=MODEL_DIRS["anydoc"],
                source_record=record,
                source_text=source_text,
                results_root=results_root,
                index=indexes["anydoc"],
            ),
            "docling": docling_output_metrics(
                source_record=record,
                source_text=source_text,
                results_root=results_root,
                index=indexes["docling"],
            ),
            "markitdown": manifest_output_metrics(
                parser_id="markitdown",
                model_dir=MODEL_DIRS["markitdown"],
                source_record=record,
                source_text=source_text,
                results_root=results_root,
                index=indexes["markitdown"],
            ),
            "mineru": mineru_output_metrics(
                source_record=record,
                source_text=source_text,
                results_root=results_root,
            ),
        }
        record["scenarios"] = derive_scenarios(record)
        record["issues"] = {
            parser: issue_attribution(record, output)
            for parser, output in record["outputs"].items()
        }
        record["audit_signals"] = {
            parser: audit_signals(record, output)
            for parser, output in record["outputs"].items()
        }
    return records


def derive_scenarios(record: dict[str, Any]) -> list[str]:
    preflight = record["preflight"]
    outputs = record.get("outputs", {})
    scenarios: set[str] = {f"format_{record['extension'].lstrip('.')}"}
    page_count = preflight.get("page_count")
    scanned_ratio = preflight.get("scanned_page_ratio")
    if record["extension"] == ".pdf":
        if isinstance(scanned_ratio, (int, float)) and scanned_ratio >= 0.8:
            scenarios.add("scanned_ocr_pdf")
        elif isinstance(scanned_ratio, (int, float)) and scanned_ratio <= 0.2:
            scenarios.add("ordinary_text_pdf")
        else:
            scenarios.add("mixed_text_scan_pdf")
    if (
        (preflight.get("source_image_object_count") or 0) >= 5
        or (preflight.get("source_media_count") or 0) > 0
        or record["category"] == "图文回答数据集源文件"
    ):
        scenarios.add("image_rich")
    source_tables = preflight.get("source_table_count") or 0
    source_cells = preflight.get("source_cell_count") or 0
    detected_tables = sorted(
        (output.get("native_table_count") or 0) for output in outputs.values()
    )
    if (
        source_tables >= 3
        or source_cells >= 100
        or (len(detected_tables) >= 2 and detected_tables[-2] >= 3)
    ):
        scenarios.add("complex_table")
    if record["extension"] in {".doc", ".ppt", ".xls"}:
        scenarios.add("legacy_office")
    if isinstance(page_count, int) and page_count >= 200:
        scenarios.add("very_long_document")
    if any(output.get("split_document") for output in outputs.values()):
        scenarios.add("split_document")
    if any(not output.get("parser_success") for output in outputs.values()):
        scenarios.add("parser_failure")
    if any(
        output.get("parser_success") and output.get("is_empty")
        for output in outputs.values()
    ):
        scenarios.add("empty_pseudo_success")
    if any((output.get("data_uri_count") or 0) > 0 for output in outputs.values()):
        scenarios.add("data_uri")
    if any(output.get("has_giant_block") for output in outputs.values()):
        scenarios.add("giant_block")
    return sorted(scenarios)


def issue_attribution(
    source_record: dict[str, Any], output: dict[str, Any]
) -> list[dict[str, Any]]:
    parser = output["parser"]
    issues: list[dict[str, Any]] = []
    evidence_base = {
        "source_path": source_record["source_path"],
        "result_directory": output.get("result_directory"),
    }
    if not output.get("parser_success"):
        issues.append(
            {
                "type": "parser_problem",
                "code": "parser_failed",
                "evidence": {
                    **evidence_base,
                    "status": output.get("parser_status"),
                    "error_code": output.get("error_code"),
                    "error": output.get("error"),
                },
            }
        )
        return issues
    if output.get("is_empty"):
        issues.append(
            {
                "type": "parser_problem",
                "code": "successful_status_but_empty_content",
                "evidence": {
                    **evidence_base,
                    "status": output.get("parser_status"),
                    "markdown_chars": 0,
                },
            }
        )
    if (
        (output.get("native_table_count") or 0) > 0
        and output.get("html_table_count", 0) == 0
        and output.get("pipe_table_row_count", 0) == 0
    ):
        issues.append(
            {
                "type": "normalization_problem",
                "code": "native_tables_not_structured_in_markdown_handoff",
                "evidence": {
                    **evidence_base,
                    "native_table_count": output.get("native_table_count"),
                    "native_table_cell_count": output.get("native_table_cell_count"),
                    "html_table_count": output.get("html_table_count"),
                    "pipe_table_row_count": output.get("pipe_table_row_count"),
                },
            }
        )
    if (
        (output.get("native_picture_count") or 0) > 0
        and output.get("image_reference_count", 0) == 0
    ):
        issues.append(
            {
                "type": "normalization_problem",
                "code": "native_pictures_not_referenced_in_markdown_handoff",
                "evidence": {
                    **evidence_base,
                    "native_picture_count": output.get("native_picture_count"),
                    "actual_asset_count": output.get("actual_asset_count"),
                    "image_reference_count": output.get("image_reference_count"),
                },
            }
        )
    if output.get("split_document") and not output.get("root_structure_merged"):
        issues.append(
            {
                "type": "normalization_problem",
                "code": "split_parts_lack_root_merged_structure",
                "evidence": {
                    **evidence_base,
                    "parts_expected": output.get("parts_expected"),
                    "parts_actual": output.get("parts_actual"),
                    "parts_complete": output.get("parts_complete"),
                    "page_ranges_contiguous": output.get("page_ranges_contiguous"),
                    "root_markdown_is_index": output.get("root_markdown_is_index"),
                },
            }
        )
    if output.get("data_uri_count", 0) > 0:
        issues.append(
            {
                "type": "quality_rule_gap",
                "code": "embedded_data_uri_requires_externalization_or_rejection",
                "evidence": {
                    **evidence_base,
                    "data_uri_count": output.get("data_uri_count"),
                    "max_block_chars": output.get("max_block_chars"),
                },
                "scope": (
                    "A required rule for the redesigned quality layer; this record "
                    "does not prove the current layer passed it."
                ),
            }
        )
    if output.get("missing_image_reference_count", 0) > 0:
        issues.append(
            {
                "type": "quality_rule_gap",
                "code": "missing_local_asset_reference_requires_blocking_rule",
                "evidence": {
                    **evidence_base,
                    "missing_reference_count": output.get(
                        "missing_image_reference_count"
                    ),
                },
                "scope": (
                    "A required rule for the redesigned quality layer; this record "
                    "does not prove the current layer passed it."
                ),
            }
        )
    return issues


def audit_signals(
    source_record: dict[str, Any], output: dict[str, Any]
) -> list[dict[str, Any]]:
    signals: list[dict[str, Any]] = []
    proxy = output.get("text_fidelity_proxy") or {}
    if (
        proxy.get("available")
        and isinstance(proxy.get("anchor_coverage"), (int, float))
        and proxy["anchor_coverage"] < 0.5
        and not output.get("is_empty")
    ):
        signals.append(
            {
                "code": "low_exact_anchor_coverage_proxy",
                "value": proxy["anchor_coverage"],
                "method": proxy.get("method"),
                "interpretation": (
                    "Screening signal only. It is not semantic accuracy and can be "
                    "reduced by OCR, formatting changes, or a defective source text layer."
                ),
            }
        )
    source_pages = source_record["preflight"].get("page_count")
    coverage = output.get("page_coverage_ratio")
    if (
        source_record["extension"] == ".pdf"
        and isinstance(source_pages, int)
        and source_pages > 1
        and isinstance(coverage, (int, float))
        and coverage < 0.8
    ):
        signals.append(
            {
                "code": "low_native_page_evidence_coverage",
                "source_pages": source_pages,
                "covered_pages": output.get("native_page_coverage_count"),
                "value": coverage,
                "interpretation": (
                    "Screening signal only. A page without emitted blocks is not "
                    "necessarily a lost page."
                ),
            }
        )
    return signals


def rating(
    label: str,
    rationale: str,
    evidence: dict[str, Any],
    *,
    method: str = "automated_evidence_triage",
) -> dict[str, Any]:
    return {
        "rating": label,
        "method": method,
        "rationale": rationale,
        "evidence": evidence,
    }


def capability_ratings(
    source_record: dict[str, Any], output: dict[str, Any]
) -> dict[str, dict[str, Any]]:
    source_evidence = {
        "source_path": source_record["source_path"],
        "page_count": source_record["preflight"].get("page_count"),
        "sheet_count": source_record["preflight"].get("sheet_count"),
    }
    output_evidence = {
        "result_directory": output.get("result_directory"),
        "parser_status": output.get("parser_status"),
        "markdown_chars": output.get("markdown_chars"),
    }
    common = {**source_evidence, **output_evidence}
    if not output.get("parser_success") or output.get("is_empty"):
        failed_evidence = {
            **common,
            "error_code": output.get("error_code"),
            "error": output.get("error"),
        }
        return {
            "text_fidelity": rating(
                "failed", "Parser failed or emitted no consumable content.", failed_evidence
            ),
            "reading_order": rating(
                "failed", "Parser failed or emitted no consumable content.", failed_evidence
            ),
            "page_traceability": rating(
                "failed" if source_record["extension"] == ".pdf" else "not_applicable",
                "No consumable output exists for page traceability."
                if source_record["extension"] == ".pdf"
                else "Page traceability is not rated for this failed non-PDF result.",
                failed_evidence,
            ),
            "table_fidelity": rating(
                "failed" if "complex_table" in source_record["scenarios"] else "not_applicable",
                "No consumable table output exists."
                if "complex_table" in source_record["scenarios"]
                else "The source is not marked table-heavy.",
                failed_evidence,
            ),
            "image_fidelity": rating(
                "failed" if "image_rich" in source_record["scenarios"] else "not_applicable",
                "No consumable image output exists."
                if "image_rich" in source_record["scenarios"]
                else "The source is not marked image-rich.",
                failed_evidence,
            ),
            "heading_structure": rating(
                "failed", "Parser failed or emitted no consumable content.", failed_evidence
            ),
            "ocr_quality": rating(
                "failed" if "scanned_ocr_pdf" in source_record["scenarios"] else "not_applicable",
                "Scan-dominant source produced no OCR content."
                if "scanned_ocr_pdf" in source_record["scenarios"]
                else "Source is not classified as a scan-dominant PDF.",
                failed_evidence,
            ),
        }

    proxy = output.get("text_fidelity_proxy") or {}
    if not proxy.get("available"):
        text_rating = rating(
            "not_applicable",
            "No independent extractable source text is available for automated comparison.",
            {**common, "proxy_reason": proxy.get("reason")},
        )
        order_rating = rating(
            "not_applicable",
            "Reading order cannot be scored without independent source anchors.",
            {**common, "proxy_reason": proxy.get("reason")},
        )
    else:
        coverage = proxy.get("anchor_coverage")
        length_ratio = proxy.get("length_ratio")
        if (
            isinstance(coverage, (int, float))
            and coverage >= 0.8
            and isinstance(length_ratio, (int, float))
            and 0.6 <= length_ratio <= 1.8
        ):
            text_label = "good"
        elif isinstance(coverage, (int, float)) and coverage >= 0.3:
            text_label = "partial"
        else:
            text_label = "failed"
        text_rating = rating(
            text_label,
            (
                "Exact normalized anchors and length ratio indicate coverage only; "
                "they do not prove semantic or OCR accuracy."
            ),
            {
                **common,
                "anchor_coverage": coverage,
                "matched_anchor_count": proxy.get("matched_anchor_count"),
                "anchor_count": proxy.get("anchor_count"),
                "length_ratio": length_ratio,
            },
        )
        order_score = proxy.get("anchor_order_score")
        if order_score is None:
            order_label = "not_applicable"
        elif order_score >= 0.95:
            order_label = "good"
        elif order_score >= 0.6:
            order_label = "partial"
        else:
            order_label = "failed"
        order_rating = rating(
            order_label,
            "Monotonicity of matched source anchors is a reading-order proxy.",
            {
                **common,
                "anchor_order_score": order_score,
                "matched_anchor_count": proxy.get("matched_anchor_count"),
            },
        )

    source_pages = source_record["preflight"].get("page_count")
    page_coverage = output.get("page_coverage_ratio")
    if source_record["extension"] != ".pdf":
        page_rating = rating(
            "not_applicable",
            "Native page coverage ratios are restricted to PDF; Office page metadata can be stale.",
            common,
        )
    elif not isinstance(source_pages, int):
        page_rating = rating(
            "not_applicable",
            "Source page count is unavailable for this format.",
            common,
        )
    elif page_coverage is None:
        page_rating = rating(
            "failed",
            "Output has no native page evidence that can be checked against the source.",
            {
                **common,
                "native_page_coverage_count": output.get("native_page_coverage_count"),
            },
        )
    else:
        page_label = (
            "good"
            if page_coverage >= 0.95
            else ("partial" if page_coverage > 0 else "failed")
        )
        page_rating = rating(
            page_label,
            "Native page evidence coverage compared with source page count.",
            {
                **common,
                "native_page_coverage_count": output.get("native_page_coverage_count"),
                "page_coverage_ratio": page_coverage,
                "page_number_examples": (output.get("native_page_numbers") or [])[:10],
            },
        )

    table_expected = "complex_table" in source_record["scenarios"]
    if not table_expected:
        table_rating = rating(
            "not_applicable",
            "Source preflight and cross-parser evidence do not mark this as table-heavy.",
            common,
        )
    else:
        table_count = output.get("native_table_count") or 0
        table_cells = output.get("native_table_cell_count") or 0
        markdown_table_evidence = (
            (output.get("html_table_count") or 0)
            + (output.get("pipe_table_row_count") or 0)
        )
        if table_cells > 0 and markdown_table_evidence > 0:
            table_label = "good"
        elif table_count > 0 or markdown_table_evidence > 0:
            table_label = "partial"
        else:
            table_label = "failed"
        table_rating = rating(
            table_label,
            (
                "Automated structure evidence cannot verify merged cells or every "
                "key value; visual/sample review remains necessary."
            ),
            {
                **common,
                "source_table_count": source_record["preflight"].get(
                    "source_table_count"
                ),
                "source_cell_count": source_record["preflight"].get("source_cell_count"),
                "native_table_count": table_count,
                "native_table_cell_count": table_cells,
                "html_table_count": output.get("html_table_count"),
                "pipe_table_row_count": output.get("pipe_table_row_count"),
            },
        )

    image_expected = "image_rich" in source_record["scenarios"]
    if not image_expected:
        image_rating = rating(
            "not_applicable",
            "Source preflight does not mark this document as image-rich.",
            common,
        )
    else:
        references = output.get("image_reference_count") or 0
        assets = output.get("actual_asset_count") or 0
        native_pictures = output.get("native_picture_count") or 0
        if references > 0 and output.get("missing_image_reference_count", 0) == 0:
            image_label = "good"
        elif references > 0 or assets > 0 or native_pictures > 0:
            image_label = "partial"
        else:
            image_label = "failed"
        image_rating = rating(
            image_label,
            "Presence and referential integrity are checked; caption binding needs visual review.",
            {
                **common,
                "source_image_object_count": source_record["preflight"].get(
                    "source_image_object_count"
                ),
                "source_media_count": source_record["preflight"].get(
                    "source_media_count"
                ),
                "native_picture_count": native_pictures,
                "image_reference_count": references,
                "actual_asset_count": assets,
                "missing_image_reference_count": output.get(
                    "missing_image_reference_count"
                ),
            },
        )

    heading_count = output.get("heading_count") or 0
    block_count = output.get("markdown_block_count") or 0
    heading_label = "good" if heading_count >= 2 else ("partial" if block_count else "failed")
    heading_rating = rating(
        heading_label,
        "Markdown heading and block presence is a structure proxy, not a hierarchy ground truth.",
        {
            **common,
            "heading_count": heading_count,
            "markdown_block_count": block_count,
        },
    )

    if "scanned_ocr_pdf" not in source_record["scenarios"]:
        ocr_rating = rating(
            "not_applicable",
            "Source is not classified as a scan-dominant PDF.",
            common,
        )
    elif output.get("markdown_chars", 0) == 0:
        ocr_rating = rating(
            "failed",
            "Scan-dominant source produced no OCR text.",
            common,
        )
    else:
        ocr_rating = rating(
            "partial",
            (
                "OCR text exists, but no character-level ground truth is available; "
                "partial is the strongest automated rating."
            ),
            {
                **common,
                "scanned_page_ratio": source_record["preflight"].get(
                    "scanned_page_ratio"
                ),
                "output_chars": output.get("markdown_chars"),
            },
        )

    return {
        "text_fidelity": text_rating,
        "reading_order": order_rating,
        "page_traceability": page_rating,
        "table_fidelity": table_rating,
        "image_fidelity": image_rating,
        "heading_structure": heading_rating,
        "ocr_quality": ocr_rating,
    }


def select_samples(
    records: list[dict[str, Any]], sample_count: int
) -> list[tuple[dict[str, Any], list[str]]]:
    minimum_count = max(24, min(32, sample_count))
    selected: dict[str, tuple[dict[str, Any], list[str]]] = {}
    selected_by_unique_source: dict[str, str] = {}

    def add(record: dict[str, Any], reason: str) -> None:
        key = record["document_key"]
        unique_source_id = record["unique_source_id"]
        existing_key = selected_by_unique_source.get(unique_source_id)
        if existing_key is not None:
            selected[existing_key][1].append(reason)
            return
        if key not in selected:
            selected[key] = (record, [])
            selected_by_unique_source[unique_source_id] = key
        selected[key][1].append(reason)

    buckets: list[tuple[str, int]] = [
        ("format_doc", 2),
        ("format_docx", 2),
        ("format_ppt", 2),
        ("format_xls", 2),
        ("format_xlsx", 1),
        ("ordinary_text_pdf", 3),
        ("scanned_ocr_pdf", 3),
        ("image_rich", 3),
        ("complex_table", 3),
        ("split_document", 2),
        ("parser_failure", 3),
        ("empty_pseudo_success", 3),
        ("data_uri", 2),
    ]
    for scenario, limit in buckets:
        candidates = sorted(
            (record for record in records if scenario in record["scenarios"]),
            key=lambda record: (
                -len(record["scenarios"]),
                -sum(len(items) for items in record["issues"].values()),
                record["document_key"].casefold(),
            ),
        )
        for record in candidates[:limit]:
            add(record, scenario)

    ranked = sorted(
        records,
        key=lambda record: (
            -sum(len(items) for items in record["issues"].values()),
            -len(record["scenarios"]),
            record["document_key"].casefold(),
        ),
    )
    for record in ranked:
        if len(selected) >= minimum_count:
            break
        add(record, "anomaly_diversity_fill")

    if len(selected) > 32:
        required_formats = {
            "format_doc",
            "format_docx",
            "format_ppt",
            "format_xls",
            "format_xlsx",
        }
        ordered = sorted(
            selected.values(),
            key=lambda pair: (
                not bool(required_formats.intersection(pair[1])),
                -len(pair[1]),
                -sum(len(items) for items in pair[0]["issues"].values()),
                pair[0]["document_key"].casefold(),
            ),
        )
        selected = {
            pair[0]["document_key"]: pair for pair in ordered[:32]
        }
    return sorted(selected.values(), key=lambda pair: pair[0]["document_key"].casefold())


def load_manual_overrides(path: Path) -> dict[tuple[str, str, str], dict[str, Any]]:
    overrides: dict[tuple[str, str, str], dict[str, Any]] = {}
    if not path.is_file():
        return overrides
    for line in read_text(path).splitlines():
        if not line.strip():
            continue
        item = json.loads(line)
        key = (item["document_key"], item["parser"], item["capability"])
        overrides[key] = item
    return overrides


def build_sample_audit(
    samples: list[tuple[dict[str, Any], list[str]]],
    overrides: dict[tuple[str, str, str], dict[str, Any]],
) -> list[dict[str, Any]]:
    audits: list[dict[str, Any]] = []
    for record, reasons in samples:
        outputs: dict[str, Any] = {}
        for parser, output in record["outputs"].items():
            capabilities = capability_ratings(record, output)
            for capability, value in list(capabilities.items()):
                override = overrides.get((record["document_key"], parser, capability))
                if override:
                    capabilities[capability] = {
                        "rating": override["rating"],
                        "method": "manual_source_comparison",
                        "rationale": override["rationale"],
                        "evidence": override["evidence"],
                    }
            outputs[parser] = {
                "capabilities": capabilities,
                "issues": record["issues"][parser],
                "audit_signals": record.get("audit_signals", {}).get(parser, []),
            }
        audits.append(
            {
                "document_key": record["document_key"],
                "document_id": record["document_id"],
                "unique_source_id": record["unique_source_id"],
                "source_path": record["source_path"],
                "category": record["category"],
                "filename": record["filename"],
                "extension": record["extension"],
                "scenarios": record["scenarios"],
                "selection_reasons": sorted(set(reasons)),
                "source_preflight": record["preflight"],
                "outputs": outputs,
                "audit_scope": (
                    "Ratings are evidence triage unless method is manual_source_comparison. "
                    "They are not absolute accuracy labels."
                ),
            }
        )
    return audits


def numeric_summary(values: Iterable[Any]) -> dict[str, Any]:
    numbers = sorted(
        float(value)
        for value in values
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    )
    if not numbers:
        return {"count": 0}

    def percentile(fraction: float) -> float:
        if len(numbers) == 1:
            return numbers[0]
        position = (len(numbers) - 1) * fraction
        lower = int(position)
        upper = min(lower + 1, len(numbers) - 1)
        weight = position - lower
        return numbers[lower] * (1 - weight) + numbers[upper] * weight

    return {
        "count": len(numbers),
        "min": round(numbers[0], 6),
        "p25": round(percentile(0.25), 6),
        "median": round(statistics.median(numbers), 6),
        "p75": round(percentile(0.75), 6),
        "max": round(numbers[-1], 6),
        "mean": round(statistics.fmean(numbers), 6),
    }


def alignment_inventory(
    source_records: list[dict[str, Any]], results_root: Path
) -> dict[str, Any]:
    expected_names = {
        f"{record['category']}/{record['filename']}".lower()
        for record in source_records
    }
    inventory: dict[str, Any] = {}
    for parser in ("anydoc", "docling", "markitdown"):
        actual_names = set(
            manifest_index(results_root, MODEL_DIRS[parser]).keys()
        )
        inventory[parser] = {
            "expected_occurrence_count": len(expected_names),
            "manifest_occurrence_count": len(actual_names),
            "missing_manifest_keys": sorted(expected_names - actual_names),
            "extra_manifest_keys": sorted(actual_names - expected_names),
            "alignment_explanation": (
                "Parser status=error remains an aligned occurrence; only absent keys "
                "appear in missing_manifest_keys."
            ),
        }

    expected_mineru = {
        (
            f"{canonical_windows_component(record['category'])}/"
            f"{canonical_windows_component(Path(record['filename']).stem)}"
        ): f"{record['category']}/{Path(record['filename']).stem}"
        for record in source_records
    }
    actual_mineru: dict[str, str] = {}
    ignored_auxiliary: list[str] = []
    for category in CATEGORIES:
        category_dir = results_root / MODEL_DIRS["mineru"] / category
        if category_dir.is_dir():
            for path in category_dir.iterdir():
                if not path.is_dir():
                    continue
                display_key = f"{category}/{path.name}"
                if path.name.startswith("_"):
                    ignored_auxiliary.append(display_key)
                    continue
                canonical_key = (
                    f"{canonical_windows_component(category)}/"
                    f"{canonical_windows_component(path.name)}"
                )
                actual_mineru[canonical_key] = display_key
    missing_keys = sorted(set(expected_mineru) - set(actual_mineru))
    extra_keys = sorted(set(actual_mineru) - set(expected_mineru))
    inventory["mineru"] = {
        "expected_occurrence_count": len(source_records),
        "expected_unique_result_key_count": len(expected_mineru),
        "result_directory_count": len(actual_mineru),
        "missing_result_keys": [expected_mineru[key] for key in missing_keys],
        "extra_result_keys": [actual_mineru[key] for key in extra_keys],
        "ignored_auxiliary_directories": sorted(ignored_auxiliary),
        "alignment_explanation": (
            "MinerU has no category manifest; category plus Windows-normalized source "
            "stem is used for directory alignment. Trailing spaces/dots are normalized, "
            "and underscore-prefixed batch-control directories are reported separately "
            "rather than counted as parser outputs."
        ),
    }
    return inventory


def model_metric_totals(outputs: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "markdown_chars": sum(output.get("markdown_chars") or 0 for output in outputs),
        "markdown_blocks": sum(
            output.get("markdown_block_count") or 0 for output in outputs
        ),
        "headings": sum(output.get("heading_count") or 0 for output in outputs),
        "native_tables": sum(
            output.get("native_table_count") or 0 for output in outputs
        ),
        "native_table_cells": sum(
            output.get("native_table_cell_count") or 0 for output in outputs
        ),
        "html_tables": sum(
            output.get("html_table_count") or 0 for output in outputs
        ),
        "image_references": sum(
            output.get("image_reference_count") or 0 for output in outputs
        ),
        "actual_assets": sum(
            output.get("actual_asset_count") or 0 for output in outputs
        ),
        "missing_image_references": sum(
            output.get("missing_image_reference_count") or 0 for output in outputs
        ),
        "data_uris": sum(output.get("data_uri_count") or 0 for output in outputs),
        "duplicate_blocks": sum(
            output.get("duplicate_block_count") or 0 for output in outputs
        ),
    }


def model_group_summary(
    records: list[dict[str, Any]], parser: str
) -> dict[str, Any]:
    outputs = [record["outputs"][parser] for record in records]
    return {
        "document_occurrences": len(records),
        "parser_success_count": sum(
            bool(output.get("parser_success")) for output in outputs
        ),
        "nonempty_markdown_count": sum(
            bool(output.get("markdown_chars")) for output in outputs
        ),
        "empty_success_count": sum(
            bool(output.get("parser_success")) and bool(output.get("is_empty"))
            for output in outputs
        ),
        "native_table_count": sum(
            output.get("native_table_count") or 0 for output in outputs
        ),
        "native_table_cell_count": sum(
            output.get("native_table_cell_count") or 0 for output in outputs
        ),
        "image_reference_count": sum(
            output.get("image_reference_count") or 0 for output in outputs
        ),
        "page_coverage_ratio": numeric_summary(
            output.get("page_coverage_ratio") for output in outputs
        ),
        "anchor_coverage": numeric_summary(
            (output.get("text_fidelity_proxy") or {}).get("anchor_coverage")
            for output in outputs
        ),
    }


def summarize_model(
    records: list[dict[str, Any]], parser: str
) -> dict[str, Any]:
    outputs = [record["outputs"][parser] for record in records]
    first_by_hash: dict[str, dict[str, Any]] = {}
    groups_by_hash: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        first_by_hash.setdefault(record["sha256"], record)
        groups_by_hash[record["sha256"]].append(record)
    unique_outputs = [
        record["outputs"][parser] for record in first_by_hash.values()
    ]
    repeated_groups = [group for group in groups_by_hash.values() if len(group) > 1]
    consistent_markdown_groups = 0
    inconsistent_markdown_groups: list[list[str]] = []
    status_variation_groups: list[list[str]] = []
    for group in repeated_groups:
        markdown_hashes = {
            record["outputs"][parser].get("markdown_sha256")
            for record in group
        }
        statuses = {
            record["outputs"][parser].get("parser_status") for record in group
        }
        document_keys = sorted(record["document_key"] for record in group)
        if len(markdown_hashes) == 1:
            consistent_markdown_groups += 1
        else:
            inconsistent_markdown_groups.append(document_keys)
        if len(statuses) > 1:
            status_variation_groups.append(document_keys)
    issue_types: Counter[str] = Counter()
    issue_codes: Counter[str] = Counter()
    signal_codes: Counter[str] = Counter()
    for record in records:
        for issue in record["issues"][parser]:
            issue_types[issue["type"]] += 1
            issue_codes[f"{issue['type']}:{issue['code']}"] += 1
        for signal in record.get("audit_signals", {}).get(parser, []):
            signal_codes[signal["code"]] += 1
    return {
        "document_occurrences": len(outputs),
        "status_counts": dict(sorted(Counter(
            output.get("parser_status") or "unknown" for output in outputs
        ).items())),
        "parser_success_count": sum(bool(output.get("parser_success")) for output in outputs),
        "result_exists_count": sum(bool(output.get("result_exists")) for output in outputs),
        "nonempty_markdown_count": sum(
            bool(output.get("markdown_chars")) for output in outputs
        ),
        "empty_success_count": sum(
            bool(output.get("parser_success")) and bool(output.get("is_empty"))
            for output in outputs
        ),
        "split_document_count": sum(
            bool(output.get("split_document")) for output in outputs
        ),
        "unique_source_representative": {
            "count": len(unique_outputs),
            "parser_success_count": sum(
                bool(output.get("parser_success")) for output in unique_outputs
            ),
            "nonempty_markdown_count": sum(
                bool(output.get("markdown_chars")) for output in unique_outputs
            ),
            "totals": model_metric_totals(unique_outputs),
            "selection_rule": (
                "First category/filename occurrence in deterministic corpus order for "
                "each exact SHA-256 group."
            ),
        },
        "duplicate_output_consistency": {
            "repeated_sha256_group_count": len(repeated_groups),
            "consistent_markdown_hash_group_count": consistent_markdown_groups,
            "inconsistent_markdown_hash_group_count": len(inconsistent_markdown_groups),
            "inconsistent_markdown_document_key_groups": inconsistent_markdown_groups,
            "status_variation_group_count": len(status_variation_groups),
            "status_variation_note": (
                "Status differences such as success versus reused are reported "
                "separately and do not imply different parsed content."
            ),
            "status_variation_document_key_groups": status_variation_groups,
        },
        "totals_by_document_occurrence": model_metric_totals(outputs),
        "by_format": {
            extension: model_group_summary(
                [record for record in records if record["extension"] == extension],
                parser,
            )
            for extension in sorted({record["extension"] for record in records})
        },
        "by_scenario": {
            scenario: model_group_summary(
                [record for record in records if scenario in record["scenarios"]],
                parser,
            )
            for scenario in (
                "ordinary_text_pdf",
                "scanned_ocr_pdf",
                "image_rich",
                "complex_table",
                "legacy_office",
                "split_document",
                "data_uri",
            )
        },
        "flag_counts": {
            "giant_block": sum(
                bool(output.get("has_giant_block")) for output in outputs
            ),
            "has_duplicate_blocks": sum(
                (output.get("duplicate_block_count") or 0) > 0 for output in outputs
            ),
            "has_data_uri": sum(
                (output.get("data_uri_count") or 0) > 0 for output in outputs
            ),
            "has_missing_image_reference": sum(
                (output.get("missing_image_reference_count") or 0) > 0
                for output in outputs
            ),
        },
        "distributions": {
            "markdown_chars": numeric_summary(
                output.get("markdown_chars") for output in outputs
            ),
            "markdown_blocks": numeric_summary(
                output.get("markdown_block_count") for output in outputs
            ),
            "page_coverage_ratio": numeric_summary(
                output.get("page_coverage_ratio") for output in outputs
            ),
            "native_bbox_coverage": numeric_summary(
                output.get("native_bbox_coverage") for output in outputs
            ),
            "native_source_id_coverage": numeric_summary(
                output.get("native_source_id_coverage") for output in outputs
            ),
            "anchor_coverage": numeric_summary(
                (output.get("text_fidelity_proxy") or {}).get("anchor_coverage")
                for output in outputs
            ),
            "anchor_order_score": numeric_summary(
                (output.get("text_fidelity_proxy") or {}).get("anchor_order_score")
                for output in outputs
            ),
            "length_ratio": numeric_summary(
                (output.get("text_fidelity_proxy") or {}).get("length_ratio")
                for output in outputs
            ),
        },
        "issue_type_counts": dict(sorted(issue_types.items())),
        "issue_code_counts": dict(sorted(issue_codes.items())),
        "audit_signal_counts": dict(sorted(signal_codes.items())),
        "normalized_lineage_note": (
            "Normalized source-id coverage is unavailable because persisted "
            "ParsedDocument/source-map outputs were not part of the supplied result set."
        ),
    }


def sample_rating_summary(
    sample_audits: list[dict[str, Any]]
) -> dict[str, Any]:
    counts: defaultdict[str, defaultdict[str, Counter[str]]] = defaultdict(
        lambda: defaultdict(Counter)
    )
    methods: Counter[str] = Counter()
    for audit in sample_audits:
        for parser, parser_audit in audit["outputs"].items():
            for capability, result in parser_audit["capabilities"].items():
                counts[parser][capability][result["rating"]] += 1
                methods[result["method"]] += 1
    return {
        "sample_document_count": len(sample_audits),
        "rating_counts": {
            parser: {
                capability: dict(sorted(ratings.items()))
                for capability, ratings in sorted(capabilities.items())
            }
            for parser, capabilities in sorted(counts.items())
        },
        "method_counts": dict(sorted(methods.items())),
    }


def build_summary(
    records: list[dict[str, Any]],
    samples: list[dict[str, Any]],
    source_root: Path,
    results_root: Path,
) -> dict[str, Any]:
    hash_counts = Counter(record["sha256"] for record in records)
    unique_sources = len(hash_counts)
    repeated_hashes = {digest: count for digest, count in hash_counts.items() if count > 1}
    global_issue_types: Counter[str] = Counter()
    global_issue_codes: Counter[str] = Counter()
    global_signal_codes: Counter[str] = Counter()
    for record in records:
        for issues in record["issues"].values():
            for issue in issues:
                global_issue_types[issue["type"]] += 1
                global_issue_codes[f"{issue['type']}:{issue['code']}"] += 1
        for signals in record.get("audit_signals", {}).values():
            for signal in signals:
                global_signal_codes[signal["code"]] += 1
    preflights = [record["preflight"] for record in records]
    return {
        "schema": "source-output-fidelity-baseline",
        "methodology": {
            "source_root": str(source_root),
            "results_root": str(results_root),
            "read_only_inputs": True,
            "parser_reruns": False,
            "primary_key": "dataset category + original filename",
            "unique_source_identity": "exact SHA-256",
            "text_comparison": (
                "Exact normalized source-text anchors are coverage/order proxies, "
                "not absolute accuracy."
            ),
            "sample_rating_scale": ["good", "partial", "failed", "not_applicable"],
            "limitations": [
                "No human-authored character, table-cell, reading-order, or OCR ground truth.",
                "Legacy OLE Office preflight exposes container metadata but not full text.",
                "Normalized ParsedDocument outputs were not persisted with the four raw result sets.",
                "Cross-parser agreement is used only to select suspicious/table-heavy samples, not as truth.",
            ],
        },
        "source_corpus": {
            "document_occurrence_count": len(records),
            "unique_source_count_by_sha256": unique_sources,
            "duplicate_occurrence_count": len(records) - unique_sources,
            "repeated_sha256_group_count": len(repeated_hashes),
            "category_counts": dict(sorted(Counter(
                record["category"] for record in records
            ).items())),
            "format_counts": dict(sorted(Counter(
                record["extension"] for record in records
            ).items())),
            "scenario_counts": dict(sorted(Counter(
                scenario for record in records for scenario in record["scenarios"]
            ).items())),
            "preflight_counts": {
                "corrupt": sum(bool(item.get("is_corrupt")) for item in preflights),
                "pdf_with_text_layer": sum(
                    item.get("source_kind") == "pdf" and item.get("has_text_layer") is True
                    for item in preflights
                ),
                "scan_dominant_pdf": sum(
                    isinstance(item.get("scanned_page_ratio"), (int, float))
                    and item["scanned_page_ratio"] >= 0.8
                    for item in preflights
                ),
                "office_with_media": sum(
                    (item.get("source_media_count") or 0) > 0 for item in preflights
                ),
                "office_with_tables": sum(
                    (item.get("source_table_count") or 0) > 0 for item in preflights
                ),
            },
            "page_count_distribution": numeric_summary(
                item.get("page_count") for item in preflights
            ),
            "sheet_count_distribution": numeric_summary(
                item.get("sheet_count") for item in preflights
            ),
        },
        "alignment": alignment_inventory(records, results_root),
        "models": {
            parser: summarize_model(records, parser)
            for parser in MODEL_DIRS
        },
        "issue_attribution": {
            "unit": "document-parser issue records; one document can contribute multiple issues",
            "type_counts": dict(sorted(global_issue_types.items())),
            "code_counts": dict(sorted(global_issue_codes.items())),
            "quality_rule_scope_note": (
                "quality_rule_gap records state rules required by the redesign; they "
                "do not by themselves prove that the current quality run passed bad content."
            ),
        },
        "audit_signals": {
            "code_counts": dict(sorted(global_signal_codes.items())),
            "interpretation": (
                "Signals identify records for review and are excluded from parser, "
                "normalization, and quality-rule problem counts."
            ),
        },
        "sample_audit": sample_rating_summary(samples),
    }


def source_manifest_projection(record: dict[str, Any]) -> dict[str, Any]:
    return {
        field: record[field]
        for field in (
            "document_key",
            "document_id",
            "unique_source_id",
            "category",
            "filename",
            "extension",
            "source_path",
            "size_bytes",
            "sha256",
            "duplicate_occurrence_count",
            "duplicate_categories",
            "duplicate_document_keys",
            "preflight",
            "scenarios",
        )
    }


def validate_results(
    records: list[dict[str, Any]],
    sample_audits: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    if len(records) != 224:
        raise RuntimeError(f"Expected 224 source occurrences, found {len(records)}")
    if not 24 <= len(sample_audits) <= 32:
        raise RuntimeError(
            f"Expected 24-32 sampled documents, found {len(sample_audits)}"
        )
    for parser, alignment in summary["alignment"].items():
        missing = alignment.get("missing_manifest_keys", alignment.get("missing_result_keys"))
        extras = alignment.get("extra_manifest_keys", alignment.get("extra_result_keys"))
        if missing or extras:
            raise RuntimeError(
                f"Unexplained alignment difference for {parser}: "
                f"missing={len(missing or [])}, extra={len(extras or [])}"
            )
    required_scenarios = {
        "format_doc",
        "format_docx",
        "format_ppt",
        "format_xls",
        "format_xlsx",
        "ordinary_text_pdf",
        "scanned_ocr_pdf",
        "image_rich",
        "complex_table",
        "split_document",
        "parser_failure",
        "empty_pseudo_success",
    }
    sampled_scenarios = {
        scenario for audit in sample_audits for scenario in audit["scenarios"]
    }
    missing_scenarios = sorted(required_scenarios - sampled_scenarios)
    if missing_scenarios:
        raise RuntimeError(f"Sample coverage missing scenarios: {missing_scenarios}")


def refresh_loaded_records(records: list[dict[str, Any]]) -> None:
    for record in records:
        if record["extension"] != ".pdf":
            for output in record["outputs"].values():
                output["page_coverage_ratio"] = None
        record["scenarios"] = derive_scenarios(record)
        record["issues"] = {
            parser: issue_attribution(record, output)
            for parser, output in record["outputs"].items()
        }
        record["audit_signals"] = {
            parser: audit_signals(record, output)
            for parser, output in record["outputs"].items()
        }


def main() -> int:
    args = parse_args()
    source_root = args.source_root.resolve()
    results_root = args.results_root.resolve()
    output_dir = args.output_dir.resolve()
    if not source_root.is_dir():
        raise FileNotFoundError(source_root)
    if not results_root.is_dir():
        raise FileNotFoundError(results_root)
    output_dir.mkdir(parents=True, exist_ok=True)

    metrics_path = output_dir / "document_metrics.jsonl"
    if args.reuse_document_metrics:
        if not metrics_path.is_file():
            raise FileNotFoundError(
                f"--reuse-document-metrics requires {metrics_path}"
            )
        records = read_jsonl(metrics_path)
        refresh_loaded_records(records)
    else:
        records, source_text_by_key = build_source_records(source_root)
        attach_outputs(records, source_text_by_key, results_root)
    selected = select_samples(records, args.sample_count)
    overrides = load_manual_overrides(output_dir / "manual_review.jsonl")
    sample_audits = build_sample_audit(selected, overrides)
    summary = build_summary(records, sample_audits, source_root, results_root)
    validate_results(records, sample_audits, summary)

    write_jsonl(
        output_dir / "source_manifest.jsonl",
        (source_manifest_projection(record) for record in records),
    )
    write_jsonl(metrics_path, records)
    write_jsonl(output_dir / "sample_audit.jsonl", sample_audits)
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(
        json.dumps(
            {
                "documents": len(records),
                "unique_sources": summary["source_corpus"]["unique_source_count_by_sha256"],
                "samples": len(sample_audits),
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
