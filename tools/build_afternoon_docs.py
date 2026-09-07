"""Build presentation-ready Word documents from the reviewed Markdown sources."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def set_font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def configure(document: Document, *, compact: bool) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.25 if compact else 1.8)
    section.bottom_margin = Cm(1.2 if compact else 1.8)
    section.left_margin = Cm(1.45 if compact else 2.0)
    section.right_margin = Cm(1.45 if compact else 2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.5 if compact else 10.5)
    normal.paragraph_format.space_after = Pt(2 if compact else 5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def convert(source: Path, target: Path, *, compact: bool, force_two_pages: bool = False) -> None:
    document = Document()
    configure(document, compact=compact)
    in_code = False
    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if force_two_pages and line == "---":
            document.add_page_break()
            continue
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7)
            set_font(paragraph.add_run(line[2:]), "Microsoft YaHei", 17 if compact else 18, bold=True)
        elif line.startswith("## "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            set_font(paragraph.add_run(line[3:]), "Microsoft YaHei", 11.5 if compact else 13, bold=True)
        elif line.startswith("### "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(1)
            set_font(paragraph.add_run(line[4:]), "Microsoft YaHei", 9.5 if compact else 11.5, bold=True)
        elif line.startswith("> "):
            paragraph = document.add_paragraph(line[2:])
            paragraph.paragraph_format.left_indent = Cm(0.4)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif in_code:
            paragraph = document.add_paragraph()
            set_font(paragraph.add_run(line), "Consolas", 8.0 if compact else 9.0)
        else:
            document.add_paragraph(line.replace("**", "").replace("`", ""))
    props = document.core_properties
    props.title = source.stem
    props.author = "朱恩铄、张云雅、叶耀华"
    document.save(target)


def main() -> None:
    convert(DOCS / "下午展示-两页结论.md", DOCS / "下午展示-两页结论.docx", compact=True, force_two_pages=True)
    convert(DOCS / "下午展示-发言稿.md", DOCS / "下午展示-发言稿.docx", compact=False)


if __name__ == "__main__":
    main()
